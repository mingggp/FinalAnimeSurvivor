"""
Generate report.docx for Final Anime Survivor project.
Uses python-docx. Run: python3 generate_report.py
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

FONT_NAME = "TH Sarabun New"
FONT_NAME_EN = "TH Sarabun New"

doc = Document()

# ─── Page margins ───────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.54)

# ─── Helper: set font for a run ─────────────────────────────────────────────
def set_font(run, size=16, bold=False, italic=False, color=None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold   = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    # Force Thai font in XML
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'),      FONT_NAME)
    rFonts.set(qn('w:hAnsi'),      FONT_NAME)
    rFonts.set(qn('w:cs'),         FONT_NAME)
    rFonts.set(qn('w:eastAsia'),   FONT_NAME)
    rPr.insert(0, rFonts)

def add_paragraph(text="", bold=False, italic=False, size=16,
                  align=WD_ALIGN_PARAGRAPH.LEFT, color=None, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic, color=color)
    return p

def add_heading1(text):
    p = add_paragraph(text, bold=True, size=20, space_before=12, space_after=6)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Add TOC bookmark entry style (Heading 1)
    pPr = p._p.get_or_add_pPr()
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), 'Heading1')
    pPr.insert(0, pStyle)
    return p

def add_heading2(text):
    p = add_paragraph(text, bold=True, size=18, space_before=8, space_after=4)
    pPr = p._p.get_or_add_pPr()
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), 'Heading2')
    pPr.insert(0, pStyle)
    return p

def add_heading3(text):
    p = add_paragraph(text, bold=True, size=16, space_before=6, space_after=3)
    return p

def add_body(text, italic=False, size=16):
    return add_paragraph(text, bold=False, italic=italic, size=size, space_before=0, space_after=4)

def add_bullet(text, size=16):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(3)
    run = p.add_run(text)
    set_font(run, size=size)
    return p

def add_image_placeholder(caption="[รูปภาพ]"):
    """Add a shaded rectangle as image placeholder."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'D9D9D9')
    pPr.append(shd)
    run = p.add_run(f"\n\n\n{caption}\n\n\n")
    set_font(run, size=14, italic=True, color=(80,80,80))
    cap_p = add_paragraph(caption, italic=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    return p

def add_table_row(table, cells_text, bold=False, size=15, bg_color=None):
    row = table.add_row()
    for i, text in enumerate(cells_text):
        cell = row.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        set_font(run, size=size, bold=bold)
        if bg_color:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  bg_color)
            tcPr.append(shd)
    return row

def insert_toc(doc):
    """Insert automatic Table of Contents field."""
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run2 = p.add_run()
    run2._r.append(fldChar3)
    note = add_paragraph(
        "(กด Ctrl+A แล้ว F9 ใน Microsoft Word เพื่ออัปเดตสารบัญ)",
        italic=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, color=(120,120,120)
    )
    return p

# ═══════════════════════════════════════════════════════════════════════════
#  COVER PAGE
# ═══════════════════════════════════════════════════════════════════════════
add_paragraph()
add_paragraph()
add_paragraph()

p = add_paragraph("รายงานโปรเจกต์", bold=True, size=28, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20)
add_paragraph("วิชา การเขียนโปรแกรมเชิงวัตถุ", bold=False, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph()

# Game title
p = add_paragraph("Final Anime Survivor", bold=True, size=32, align=WD_ALIGN_PARAGRAPH.CENTER, color=(31,73,125))
add_paragraph("เกมเอาตัวรอดธีมอนิเมะในสไตล์ Vampire Survivors", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
add_paragraph()
add_paragraph()

add_image_placeholder("[ภาพ Screenshot หน้าหลักของเกม]")
add_paragraph()
add_paragraph()

# Authors
tbl = doc.add_table(rows=1, cols=2)
tbl.style = 'Table Grid'
add_table_row(tbl, ["รหัสนิสิต", "ชื่อ-นามสกุล"], bold=True, size=16, bg_color="D6E4F0")
add_table_row(tbl, ["6831356521", "Popnipit Watcharapitak  (ปพนพิศิษฐ์ วัชรพิทักษ์)"], size=16)
add_table_row(tbl, ["6831362221", "Voravich Thanyavinichakul  (วรวิชญ์ ธัญวินิจกุล)"], size=16)
add_paragraph()

add_paragraph("ภาคเรียนที่ 1  ปีการศึกษา 2567–2568", size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
add_paragraph("คณะวิศวกรรมศาสตร์  จุฬาลงกรณ์มหาวิทยาลัย", size=15, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  สารบัญ (TABLE OF CONTENTS)
# ═══════════════════════════════════════════════════════════════════════════
add_paragraph("สารบัญ", bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12)
add_paragraph()
insert_toc(doc)
doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  1. บทนำ
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("1. บทนำ")
add_body(
    "โปรแกรม Final Anime Survivor เป็นวิดีโอเกมแนว Survivor แบบ Top-down "
    "ที่ได้รับแรงบันดาลใจจากเกม Vampire Survivors โดยมีธีมตัวละครและอาวุธจากอนิเมะ "
    "Jujutsu Kaisen ผู้เล่นจะต้องเอาชีวิตรอดจากคลื่นศัตรูที่เพิ่มขึ้นเรื่อย ๆ "
    "รวบรวมอาวุธและของเสริมต่าง ๆ เพื่อพัฒนาตัวละครให้แข็งแกร่งขึ้นในแต่ละรอบการเล่น"
)
add_body(
    "โปรแกรมพัฒนาด้วยภาษา Java โดยใช้ JavaFX เป็น Framework สำหรับการแสดงผลกราฟิก "
    "และระบบ UI และออกแบบโครงสร้างโปรแกรมตามหลักการ Object-Oriented Programming (OOP) "
    "ได้แก่ Inheritance, Interface, Polymorphism และ Encapsulation"
)

add_heading2("1.1 วัตถุประสงค์ของโครงงาน")
add_bullet("พัฒนาเกมแนว Survivor ที่มีความซับซ้อน มีระบบอาวุธ ของสะสม และการยกระดับตัวละคร")
add_bullet("ประยุกต์ใช้แนวคิด OOP ในการออกแบบโครงสร้างของโปรแกรมอย่างเป็นระบบ")
add_bullet("ฝึกทักษะการเขียน Unit Test ด้วย JUnit 5 เพื่อทดสอบ Logic ของโปรแกรม")
add_bullet("เรียนรู้การพัฒนา GUI ด้วย JavaFX และการจัดการ Scene/State ของเกม")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  2. รายละเอียดโปรแกรมและวิธีใช้งาน
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("2. รายละเอียดโปรแกรมและวิธีการใช้งาน")

add_heading2("2.1 ความต้องการของระบบ")
tbl2 = doc.add_table(rows=1, cols=2)
tbl2.style = 'Table Grid'
add_table_row(tbl2, ["ส่วนประกอบ", "เวอร์ชัน / รายละเอียด"], bold=True, size=15, bg_color="D6E4F0")
add_table_row(tbl2, ["Java JDK", "21 ขึ้นไป"], size=15)
add_table_row(tbl2, ["JavaFX", "24.0.2"], size=15)
add_table_row(tbl2, ["Gradle", "8.x (มี Wrapper ในโปรเจกต์)"], size=15)
add_table_row(tbl2, ["ระบบปฏิบัติการ", "Windows / macOS / Linux"], size=15)
add_paragraph()

add_heading2("2.2 วิธีการติดตั้งและรันโปรแกรม")

add_heading3("วิธีที่ 1 — รันจาก Source Code")
add_body("เปิด Terminal ที่โฟลเดอร์โปรเจกต์แล้วพิมพ์:")
p = doc.add_paragraph()
run = p.add_run("./gradlew run")
set_font(run, size=14, bold=True)
pPr = p._p.get_or_add_pPr()
shd = OxmlElement('w:shd')
shd.set(qn('w:val'),   'clear')
shd.set(qn('w:color'), 'auto')
shd.set(qn('w:fill'),  'F2F2F2')
pPr.append(shd)
add_paragraph()

add_heading3("วิธีที่ 2 — รันจาก JAR File")
add_body("สร้างโฟลเดอร์ใหม่ แล้ววาง JAR ไว้ในโฟลเดอร์นั้นก่อนรัน:")
p2 = doc.add_paragraph()
run2 = p2.add_run("mkdir game\ncp FinalAnimeSurvivor.jar game/\ncd game/\njava -jar FinalAnimeSurvivor.jar")
set_font(run2, size=14, bold=True)
pPr2 = p2._p.get_or_add_pPr()
shd2 = OxmlElement('w:shd')
shd2.set(qn('w:val'),   'clear')
shd2.set(qn('w:color'), 'auto')
shd2.set(qn('w:fill'),  'F2F2F2')
pPr2.append(shd2)
add_paragraph()

add_heading2("2.3 หน้าเมนูหลัก (Main Menu)")
add_body("เมื่อเปิดโปรแกรมจะแสดงหน้าเมนูหลัก มีปุ่มดังนี้:")
add_bullet("Play — เข้าสู่หน้าเลือกตัวละคร")
add_bullet("Settings — ตั้งค่าเสียง BGM (เพลงพื้นหลัง) และ SFX (เสียงเอฟเฟกต์) แยกกัน")
add_bullet("Quit — ออกจากโปรแกรม")
add_image_placeholder("[ภาพ Screenshot หน้า Main Menu]")

add_heading2("2.4 หน้าเลือกตัวละคร (Character Select)")
add_body("ผู้เล่นสามารถเลือกตัวละครหนึ่งจากสองตัวก่อนเริ่มเกม:")

tbl3 = doc.add_table(rows=1, cols=3)
tbl3.style = 'Table Grid'
add_table_row(tbl3, ["ตัวละคร", "อาวุธเริ่มต้น", "จุดเด่น"], bold=True, size=15, bg_color="D6E4F0")
add_table_row(tbl3, ["Sukuna", "Cleave (ฟันดาบรอบทิศ)", "ดาเมจระยะประชิดสูง"], size=15)
add_table_row(tbl3, ["Gojo", "Infinity — Mugen (กลมล้อมรอบ)", "ดาเมจพื้นที่ต่อเนื่อง"], size=15)
add_paragraph()
add_image_placeholder("[ภาพ Screenshot หน้าเลือกตัวละคร]")

add_heading2("2.5 การควบคุมในเกม")
tbl4 = doc.add_table(rows=1, cols=2)
tbl4.style = 'Table Grid'
add_table_row(tbl4, ["ปุ่ม", "การกระทำ"], bold=True, size=15, bg_color="D6E4F0")
add_table_row(tbl4, ["W / A / S / D", "เคลื่อนที่ตัวละคร"], size=15)
add_table_row(tbl4, ["Tab", "เปิด/ปิดกระเป๋าสัมภาระ (Backpack)"], size=15)
add_table_row(tbl4, ["Q", "เปิด/ปิดระบบใช้ของอัตโนมัติ"], size=15)
add_table_row(tbl4, ["Esc", "เปิดเมนูหยุดเกมชั่วคราว (Pause) / ออกจากหน้าเมนู"], size=15)
add_table_row(tbl4, ["W / S (ในเมนู Level Up)", "เลือกตัวเลือก"], size=15)
add_table_row(tbl4, ["Space (ในเมนู Level Up)", "ยืนยันตัวเลือก"], size=15)
add_table_row(tbl4, ["คลิกเมาส์ (ในเมนู Level Up)", "เลือกตัวเลือก"], size=15)
add_paragraph()

add_heading2("2.6 ระบบเกมหลัก")

add_heading3("ระบบ XP และการเลื่อนระดับ")
add_body(
    "ศัตรูที่ตายมีโอกาส 80% ดรอป Exp Orb สีแดง ผู้เล่นเก็บได้เมื่อวิ่งเข้าไปหรืออยู่ในรัศมี Magnet "
    "เมื่อ XP เต็มแถบจะขึ้นเมนู Level Up ให้เลือก 1 ใน 3 ตัวเลือก ได้แก่ "
    "อาวุธใหม่ อัปเกรดอาวุธ หรือ Accessory ใหม่/อัปเกรด"
)
add_image_placeholder("[ภาพ Screenshot ระบบ Level Up]")

add_heading3("ระบบอาวุธ (Weapons)")
tbl5 = doc.add_table(rows=1, cols=3)
tbl5.style = 'Table Grid'
add_table_row(tbl5, ["อาวุธ", "ประเภท", "คำอธิบาย"], bold=True, size=14, bg_color="D6E4F0")
add_table_row(tbl5, ["OSU! Cursor (Standard)", "กระสุนเด้ง", "Cursor เด้งไปมาในหน้าจอ ดาเมจศัตรูที่โดน"], size=14)
add_table_row(tbl5, ["Lapse Blue", "วงกลมเล็งเป้า", "ดรอปวงดาเมจลงบนศัตรูที่ใกล้ที่สุด"], size=14)
add_table_row(tbl5, ["Reversal Red", "กระสุน", "ยิงกระสุนที่ระเบิดเมื่อกระทบหรือถึงระยะ"], size=14)
add_table_row(tbl5, ["Cleave", "ฟันระยะประชิด", "ฟันศัตรูที่ใกล้ที่สุดอยู่ตลอดเวลา"], size=14)
add_table_row(tbl5, ["Infinity (Mugen)", "Aura รอบตัว", "วงกลมดาเมจล้อมรอบผู้เล่นตลอดเวลา"], size=14)
add_table_row(tbl5, ["Dismantle", "กระสุนเจาะ", "กระสุนทะลุผ่านศัตรูหลายตัวได้"], size=14)
add_table_row(tbl5, ["Bible", "Orbital Books", "หนังสือโคจรรอบผู้เล่น ดาเมจศัตรูที่โดนวงโคจร"], size=14)
add_paragraph()
add_image_placeholder("[ภาพ Screenshot ระบบอาวุธและ Accessory Slot]")

add_heading3("ระบบ Chest และการ Evolve")
add_body(
    "ศัตรูที่ตายมีโอกาส 0.1% ดรอป Chest สีทอง เมื่อผู้เล่นวิ่งเข้าหา Chest จะแสดงตัวเลือก "
    "เช่น Evolve อาวุธ, Unite อาวุธ หรือ Craft ไอเทม อาวุธที่ Evolve แล้วจะแข็งแกร่งกว่ามาก"
)
add_image_placeholder("[ภาพ Screenshot ระบบ Chest]")

add_heading3("ระบบกระเป๋าสัมภาระ (Backpack)")
add_body(
    "กระเป๋าสัมภาระขนาด 6×6 ช่อง ใช้เก็บของที่ดรอปจากศัตรูหรือได้รับจากระบบต่าง ๆ "
    "กดที่ไอเทมเพื่อใช้งาน (สำหรับของที่ Usable) กด Tab เพื่อเปิด/ปิด"
)
add_image_placeholder("[ภาพ Screenshot กระเป๋าสัมภาระ]")

add_heading3("หน้าหยุดเกมชั่วคราว (Pause Screen)")
add_body(
    "กด ESC ระหว่างเกมเพื่อเปิด Pause Screen ซึ่งมีปุ่มดังนี้:"
)
add_bullet("▶ Resume — กลับสู่การเล่นต่อ (หรือกด ESC อีกครั้ง)")
add_bullet("⚙ Settings — เปิดหน้าตั้งค่าเสียง BGM/SFX ได้แม้อยู่ในเกม")
add_bullet("⏹ Main Menu — ออกจากเกมและกลับสู่หน้าหลัก (เพลงจะหยุด)")
add_image_placeholder("[ภาพ Screenshot Pause Screen]")

add_heading3("หน้าตัวละครตาย (Death Screen)")
add_body(
    "เมื่อ HP ลดลงถึง 0 เกมจะแสดงหน้าผลลัพธ์บนหน้าจอ "
    "พร้อมปุ่มสองปุ่มที่ด้านล่าง:"
)
add_bullet("↺ Restart — กลับไปหน้าเลือกตัวละครเพื่อเริ่มรอบใหม่")
add_bullet("⏹ Main Menu — กลับสู่หน้าเมนูหลัก (เพลงหยุดโดยอัตโนมัติ)")
add_image_placeholder("[ภาพ Screenshot Death Screen]")

add_heading3("ระบบเสียง (Sound)")
add_body("เกมมีเสียงประกอบครบทุกการกระทำ:")
tbl_sfx = doc.add_table(rows=1, cols=2)
tbl_sfx.style = 'Table Grid'
add_table_row(tbl_sfx, ["เหตุการณ์", "เสียง"], bold=True, size=14, bg_color="D6E4F0")
add_table_row(tbl_sfx, ["กดปุ่มทุกปุ่มใน UI", "เสียง Click สั้น"], size=14)
add_table_row(tbl_sfx, ["เก็บ Exp Orb", "เสียง Sweep ขึ้น"], size=14)
add_table_row(tbl_sfx, ["Level Up", "เสียง Arpeggio C→E→G→C"], size=14)
add_table_row(tbl_sfx, ["Sukuna โจมตี (Cleave)", "เสียง Slash + Swoosh"], size=14)
add_table_row(tbl_sfx, ["Gojo โจมตี (Infinity)", "เสียง Hum + Shimmer"], size=14)
add_body("ปรับระดับเสียงได้แยกกัน ทั้ง BGM และ SFX ผ่านหน้า Settings")
add_paragraph()

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  3. การออกแบบโปรแกรม (Design)
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("3. การออกแบบโปรแกรม")

add_heading2("3.1 ภาพรวมสถาปัตยกรรม")
add_body(
    "โปรแกรมถูกแบ่งออกเป็น Package หลักดังนี้ Game Loop และ State จัดการโดย GameManager "
    "ซึ่งทำงานเป็น Central Controller อ่าน Input จาก InputManager อัปเดต Entity ทุกชนิด "
    "และส่งข้อมูลไปยัง SceneManager เพื่ออัปเดต UI ใน JavaFX"
)
add_image_placeholder("[ภาพ UML Class Diagram — ภาพรวม Package]")

add_heading2("3.2 โครงสร้าง Package")
add_bullet("application/ — Entry point: Main, Launcher")
add_bullet("core/ — GameManager, EnemySpawner, GameState")
add_bullet("entity/ — ลำดับชั้น Entity ทั้งหมด (character, enemy, weapon, accessory, item, misc, map)")
add_bullet("entityInterface/ — Interface ทุกตัวในระบบ")
add_bullet("gui/ — Canvas และ Panel ของ JavaFX")
add_bullet("tile/ — ระบบ Tile และ TileManager")
add_bullet("utils/ — CollisionChecker, InputManager, SceneManager, SoundManager, SpriteManager")
add_bullet("vfx/ — ระบบ Visual Effect (DamageText, Slash, SkillPopUp, VFXManager)")

add_heading2("3.3 UML Diagram — Entity Hierarchy")
add_image_placeholder("[ภาพ UML Class Diagram — Entity (Entity, Character, Enemy, Item, Weapon, Accessory)]")

add_heading2("3.4 UML Diagram — Interface")
add_image_placeholder("[ภาพ UML Class Diagram — entityInterface Package]")

add_heading2("3.5 UML Diagram — Core และ Utils")
add_image_placeholder("[ภาพ UML Class Diagram — GameManager, EnemySpawner, SoundManager, SceneManager]")

add_heading2("3.6 UML Diagram — GUI")
add_image_placeholder("[ภาพ UML Class Diagram — GUI Package]")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  4. แนวคิด OOP
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("4. แนวคิดเชิงวัตถุ (OOP Concepts)")

add_heading2("4.1 Inheritance (การสืบทอด)")
add_body(
    "โปรแกรมใช้ Inheritance เพื่อแบ่งปัน Attribute และ Behavior ที่เหมือนกันระหว่าง Entity "
    "ที่เกี่ยวข้องกัน โดยไม่ต้องเขียนซ้ำ ลดความซ้ำซ้อนของโค้ด และทำให้ง่ายต่อการขยาย"
)
add_heading3("ลำดับชั้น Weapon")
add_body(
    "Weapon (abstract) ← Standard, Blue, Red, Cleave, Infinity, Dismantle "
    "Weapon กำหนด Field ร่วม เช่น level, maxLevel, cooldown, icon "
    "และ Method ร่วม เช่น upgrade(), getLevel(), getCooldown() "
    "Evolved Weapon สืบทอดต่อจาก Base Weapon: MaximumBlue extends Blue เป็นต้น"
)
add_heading3("ลำดับชั้น Entity")
add_body(
    "Entity (abstract) ← Character, Enemy "
    "Entity เก็บตำแหน่งบนแผนที่ (mapX, mapY) และชื่อ "
    "Character และ Enemy ต่างขยาย Entity ด้วย HP, Speed, Hitbox และพฤติกรรมเฉพาะ"
)
add_heading3("ลำดับชั้น Item")
add_body(
    "Item (abstract) ← Soda, Harvest, SukunaFinger, SubaruShirt "
    "Item ← craftable/ ← MalevolentKitchen, UnlimitedHollowPurple "
    "Item เก็บ name, icon, amount ที่ใช้ร่วมกันทุก Item"
)
add_heading3("ลำดับชั้น Accessory")
add_body(
    "Accessory (abstract) ← SukunaArm, SukunaCloak, SixEye, GojoGlasses, Blindfold, KeyPad "
    "Accessory กำหนด level, maxLevel และ Abstract Method procEffect() "
    "ซึ่งแต่ละ Accessory Implement ต่างกันตาม Effect ของตน"
)
add_image_placeholder("[ภาพ UML แสดง Inheritance Hierarchy]")

add_heading2("4.2 Interface (ส่วนต่อประสาน)")
add_body(
    "โปรแกรมใช้ Interface เพื่อกำหนด Contract ของพฤติกรรมที่ไม่ขึ้นอยู่กับลำดับชั้น "
    "ทำให้ GameManager สามารถเรียกใช้งาน Entity ต่าง ๆ ผ่าน Interface โดยไม่ต้องรู้ Class จริง"
)

tbl_iface = doc.add_table(rows=1, cols=2)
tbl_iface.style = 'Table Grid'
add_table_row(tbl_iface, ["Interface", "วัตถุประสงค์"], bold=True, size=14, bg_color="D6E4F0")
add_table_row(tbl_iface, ["GameObject", "Base Interface ทุก Game Object ต้อง Implement"], size=14)
add_table_row(tbl_iface, ["Renderable", "Object ที่วาดบนหน้าจอได้ผ่าน render(GraphicsContext)"], size=14)
add_table_row(tbl_iface, ["Updatable", "Object ที่อัปเดตทุก Frame ผ่าน update(dt)"], size=14)
add_table_row(tbl_iface, ["Usable", "ไอเทมที่ผู้เล่นใช้งานได้ผ่าน use(slot)"], size=14)
add_table_row(tbl_iface, ["Droppable", "ไอเทมที่ดรอปและเคลื่อนที่หาผู้เล่น"], size=14)
add_table_row(tbl_iface, ["Craftable", "ไอเทมที่ Craft ได้จาก Material"], size=14)
add_table_row(tbl_iface, ["Material", "ไอเทมที่ใช้เป็นวัตถุดิบ Craft"], size=14)
add_table_row(tbl_iface, ["Unique", "ไอเทม/อาวุธที่มีได้เพียงชิ้นเดียวต่อรอบ"], size=14)
add_table_row(tbl_iface, ["Evolvable", "อาวุธที่ Evolve เป็นรูปแบบขั้นสูงได้"], size=14)
add_table_row(tbl_iface, ["Unitable", "อาวุธที่ Fuse กับอาวุธอื่นได้"], size=14)
add_table_row(tbl_iface, ["DamageIncreasable", "อาวุธที่อัปเกรดค่าดาเมจได้"], size=14)
add_table_row(tbl_iface, ["CooldownDecreasable", "อาวุธที่ลด Cooldown ได้"], size=14)
add_table_row(tbl_iface, ["SizeIncreasable", "อาวุธที่เพิ่มขนาดได้"], size=14)
add_table_row(tbl_iface, ["DurationIncreasable", "อาวุธที่เพิ่มระยะเวลาได้"], size=14)
add_table_row(tbl_iface, ["SpeedIncreasable", "อาวุธที่เพิ่มความเร็วได้"], size=14)
add_table_row(tbl_iface, ["AmountIncreasable", "อาวุธที่เพิ่มจำนวนได้"], size=14)
add_paragraph()

add_heading2("4.3 Polymorphism (พหุนิยม)")
add_body(
    "Polymorphism ถูกใช้อย่างแพร่หลายทั่วโปรแกรม เพื่อให้ GameManager จัดการ Entity "
    "หลากหลายประเภทในรูปแบบเดียวกัน:"
)
add_bullet(
    "weapon.use(dt) — Loop ผ่าน weaponList[] เรียก use() แต่ละ Weapon "
    "โดยแต่ละอาวุธมีพฤติกรรมการยิงต่างกันสมบูรณ์"
)
add_bullet(
    "droppedItem.updateAsDroppedItem(dt) — ทุก Droppable Item เคลื่อนที่หาผู้เล่นผ่าน Interface เดียวกัน"
)
add_bullet(
    "item.update(dt) / item.isExpired() — อัปเดตผ่าน Updatable Interface "
    "โดยไม่ต้องรู้ว่าเป็น Item ชนิดใด"
)
add_bullet(
    "accessory.procEffect() — แต่ละ Accessory Override procEffect() ด้วย Effect ที่ต่างกัน "
    "แต่เรียกจาก Loop เดียวกัน"
)
add_bullet(
    "gameObject.copy() — ทุก Entity Implement Prototype Pattern ผ่าน copy() "
    "เพื่อสร้าง Instance ใหม่โดยไม่รู้ Class จริง"
)

add_heading2("4.4 Access Modifier (ตัวควบคุมการเข้าถึง)")
add_body("โปรแกรมใช้ Access Modifier อย่างเหมาะสมดังนี้:")
add_bullet(
    "private — Field ทุก Field ของ Entity เช่น mapX, currentHP, damage "
    "ถูกซ่อนและเข้าถึงผ่าน Getter/Setter เพื่อป้องกันการแก้ไขโดยตรง (Encapsulation)"
)
add_bullet(
    "public — Method ที่ส่วนอื่นต้องเรียกใช้ เช่น update(), render(), use(), procEffect()"
)
add_bullet(
    "private constructor + public static getInstance() — SoundManager ใช้ Singleton Pattern "
    "ป้องกันการสร้าง Instance ซ้ำซ้อน"
)
add_bullet(
    "private — Method ภายใน GameManager เช่น getRandomWeapon(), getRandomLevelUpChoice() "
    "ไม่ถูกเปิดเผยออกภายนอก Class"
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  5. เอกสาร JavaDoc
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("5. เอกสารอธิบายโค้ด (JavaDoc)")
add_body(
    "เอกสาร JavaDoc ที่สร้างจากโค้ดโปรแกรมสามารถเข้าถึงได้ที่ลิงก์ GitHub Pages ด้านล่าง "
    "หรือดูได้จาก Folder doc/ ใน Repository:"
)
p_link = add_paragraph()
run_link = p_link.add_run("[วางลิงก์ JavaDoc ที่นี่ — เช่น https://mingggp.github.io/finalanimesurvivor/javadoc/]")
set_font(run_link, size=15, italic=True, color=(0, 70, 180))

add_body(
    "JavaDoc ครอบคลุมทุก Package หลัก ได้แก่ core, entity, entityInterface, gui, tile, utils และ vfx "
    "พร้อมคำอธิบาย Class, Field, Method และ Parameter ที่สำคัญ"
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  6. การทดสอบ (JUnit)
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("6. การทดสอบโปรแกรม (JUnit 5)")
add_body(
    "โปรแกรมมี Unit Test ที่เขียนด้วย JUnit 5 เพื่อทดสอบ Logic สำคัญของโปรแกรม "
    "Test ทั้งหมดอยู่ใน src/test/java/ และรันด้วยคำสั่ง ./gradlew test"
)

add_heading2("6.1 WeaponTest")
add_body("ทดสอบ Logic ของระบบอาวุธและ Cooldown Progress Bar:")
add_bullet("initialLevelIsOne — ตรวจสอบว่าระดับเริ่มต้นของอาวุธคือ 1")
add_bullet("upgradeIncrementsLevelUntilMax — อัปเกรดจนถึง maxLevel แล้วได้ String \"Max\"")
add_bullet("cooldownProgressIsZeroJustAfterFiring — ค่า getCooldownProgress() == 0.0 ทันทีหลังยิง")
add_bullet("cooldownProgressIsOneWhenReady — getCooldownProgress() == 1.0 เมื่อ Cooldown เต็ม")
add_bullet("cooldownProgressClampsAboveOne — ค่าไม่เกิน 1.0 แม้ timeSinceUse > cooldown")
add_bullet("cooldownProgressIsHalfWayThrough — getCooldownProgress() ≈ 0.5 ที่กึ่งกลาง")
add_bullet("zeroCooldownWeaponAlwaysReady — อาวุธที่มี cooldown = 0 พร้อมยิงเสมอ")
add_bullet("useAdvancesTimeSinceUse — use(dt) เพิ่มค่า timeSinceUse")
add_bullet("cooldownChangesAffectProgress — เปลี่ยน cooldown กระทบ progress ทันที")

add_heading2("6.2 AccessoryTest")
add_body("ทดสอบ Logic ของระบบ Accessory:")
add_bullet("initialLevelIsOne — ตรวจสอบว่าระดับเริ่มต้นของ Accessory คือ 1")
add_bullet("upgradeIncrementsLevel — upgrade() เพิ่มค่า level ได้ถูกต้อง")
add_bullet("upgradeCapsAtMaxLevel — อัปเกรดจนถึง maxLevel แล้วได้ String \"Max\"")
add_bullet("procEffectIsCallable — procEffect() เรียกได้โดยไม่ Error")
add_bullet("nameIsSetByConstructor — ชื่อที่ส่ง Constructor ถูกเก็บถูกต้อง")

add_heading2("6.3 ItemTest")
add_body("ทดสอบ Logic ของระบบ Item:")
add_bullet("defaultAmountIsZero — ค่า amount เริ่มต้นที่ 0")
add_bullet("setAmountStoresValue — setAmount() เก็บค่าได้ถูกต้อง")
add_bullet("tagFlipsMagnetFlag — tag() เปลี่ยน taggedByMagnet เป็น true")
add_bullet("nameIsSetByConstructor — ชื่อที่ส่ง Constructor ถูกเก็บถูกต้อง")

add_heading2("6.4 CollisionCheckerTest")
add_body("ทดสอบระบบตรวจสอบการชนด้วย Grid 3×3 ที่มีกำแพงล้อมรอบ:")
add_bullet("wallStopsLeftwardMovement — กำแพงทางซ้ายหยุดการเคลื่อนที่ไปซ้าย")
add_bullet("wallStopsRightwardMovement — กำแพงทางขวาหยุดการเคลื่อนที่ไปขวา")
add_bullet("wallStopsUpwardMovement — กำแพงด้านบนหยุดการเคลื่อนที่ขึ้น")
add_bullet("wallStopsDownwardMovement — กำแพงด้านล่างหยุดการเคลื่อนที่ลง")
add_bullet("noOpWhenTilesNotInitialized — ไม่ Crash เมื่อยังไม่มี Tile Map ถูกโหลด")

add_heading2("6.5 GameStateTest")
add_body("ทดสอบ Enum GameState ที่ใช้ควบคุม State Machine ของเกม:")
add_bullet("allRequiredStatesArePresent — ตรวจสอบว่ามี State ที่จำเป็น เช่น PLAYING, PAUSED, MAIN_MENU, DEATH ครบ")
add_bullet("valuesAreUnique — ทุก State มี ordinal ไม่ซ้ำกัน")
add_bullet("enumHasReasonableSize — มีอย่างน้อย 10 State (ครอบคลุมทุกโหมดของเกม)")

add_heading2("6.6 CharacterTest")
add_body("ทดสอบ Logic ของตัวละครผู้เล่น:")
add_bullet("testInitialState — ตรวจสอบค่าเริ่มต้นของตัวละคร เช่น HP, Speed")
add_bullet("testReceiveDamage — ตรวจสอบว่า HP ลดลงถูกต้องเมื่อรับดาเมจ")
add_bullet("testHeal — ตรวจสอบว่า HP ฟื้นฟูได้ถูกต้องและไม่เกิน maxHP")
add_bullet("testIsDead — ตรวจสอบว่าตัวละครตายเมื่อ HP ≤ 0")

add_heading2("6.7 EnemyTest")
add_body("ทดสอบ Logic ของศัตรู:")
add_bullet("testInitialState — ตรวจสอบค่าเริ่มต้นของศัตรู")
add_bullet("testTakeDamage — ตรวจสอบว่า HP ศัตรูลดลงถูกต้อง")
add_bullet("testMovementDirection — ตรวจสอบว่าศัตรูเคลื่อนที่เข้าหาผู้เล่นถูกทิศ")

add_heading2("6.8 ChestTest")
add_body("ทดสอบการสร้าง Chest Object:")
add_bullet("testChestCreation — ตรวจสอบว่า Chest Initialize ถูกต้อง")

add_heading2("6.9 ExpOrbTest")
add_body("ทดสอบการสร้าง ExpOrb Object:")
add_bullet("testInitialState — ตรวจสอบค่า XP Amount และ State เริ่มต้น")

add_heading2("6.10 InputManagerTest")
add_body("ทดสอบระบบรับ Input:")
add_bullet("testAddAndRemoveKey — ตรวจสอบว่า addKey/removeKey ทำงานถูกต้อง")
add_bullet("testMultipleKeys — ตรวจสอบการกดหลายปุ่มพร้อมกัน")

add_image_placeholder("[ภาพ Screenshot ผลการรัน JUnit Test (./gradlew test)]")
add_paragraph()
add_body("ตัวอย่างผลการรัน Test:", italic=True)
add_image_placeholder("[ภาพ Test Report จาก Gradle (build/reports/tests/)]")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  7. ปัญหาและข้อบกพร่อง
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("7. ปัญหาและข้อบกพร่องที่พบ")
add_body("ในระหว่างการพัฒนาพบปัญหาและข้อบกพร่องดังต่อไปนี้:")

tbl_bug = doc.add_table(rows=1, cols=4)
tbl_bug.style = 'Table Grid'
add_table_row(tbl_bug, ["#", "ส่วนที่เกิดปัญหา", "รายละเอียด", "สถานะ"], bold=True, size=14, bg_color="FFD966")

bugs = [
    ("1", "Sound (SFX)",
     "SoundManager.playSFX() ใช้ new File(path) ซึ่งไม่รองรับ Resource ภายใน JAR "
     "ทำให้เอฟเฟกต์เสียงไม่ทำงานเมื่อรันจาก JAR File "
     "แก้ไขโดยเปลี่ยนเป็น getClass().getResource(path).toExternalForm()",
     "แก้ไขแล้ว"),
    ("2", "UI — Cooldown Bar",
     "ไม่มี Visual Cooldown Bar แสดงสถานะ Cooldown ของอาวุธใน HUD "
     "WeaponSlot แสดงเพียง Level Number เท่านั้น "
     "แก้ไขโดยเพิ่ม getCooldownProgress() ใน Weapon และ Rectangle Overlay ใน WeaponSlot",
     "แก้ไขแล้ว"),
    ("3", "Shrine Icon",
     "มีไฟล์ Sprite ของ Shrine หลายเวอร์ชันใน Resources (shrine.png, SHRINE4.png, 140x140shrine.png) "
     "ภาพที่แสดงในเกมอาจไม่ตรงกับที่ตั้งใจ",
     "ยังไม่แก้ไข"),
    ("4", "Soda (ไอเทม)",
     "เงื่อนไข use() ของ Soda บังคับให้ใช้ได้เฉพาะเมื่อ HP < maxHP เท่านั้น "
     "ทำให้ไม่สามารถดื่มล่วงหน้าเมื่อเลือด HP เต็มได้",
     "ยังไม่แก้ไข"),
    ("5", "Map",
     "มีแผนที่เพียง 1 แผ่น (map1.txt) ที่เป็นสนามเล่นสี่เหลี่ยมธรรมดา "
     "ยังไม่มีความหลากหลายของ Biome หรือ Layout",
     "ยังไม่แก้ไข"),
    ("6", "NullPointerException",
     "getClosestTarget() ใน GameManager อาจ Return null เมื่อไม่มีศัตรู "
     "แต่อาวุธ Blue และ Red ไม่มีการตรวจสอบ null ก่อนใช้งาน อาจ Crash ได้ "
     "แก้ไขโดยเพิ่ม null check ก่อนเข้าถึงข้อมูลของ closestTarget",
     "แก้ไขแล้ว"),
    ("7", "SixEye Accessory",
     "procEffect() ใช้ break แทน continue เมื่อพบ weaponList slot ที่เป็น null "
     "ทำให้หยุดตรวจสอบ Slot ที่เหลือ อาวุธใน Slot ถัด ๆ ไปจะไม่ได้รับ Size Buff "
     "แก้ไขโดยเปลี่ยน break เป็น continue",
     "แก้ไขแล้ว"),
    ("8", "Bible Weapon",
     "Class Bible มีอยู่ในโปรเจกต์แต่ยังไม่ได้ Implement พฤติกรรมการโจมตี (Placeholder) "
     "แก้ไขโดย Implement ระบบ Orbital Books ที่โคจรรอบผู้เล่นและดาเมจศัตรูที่โดน",
     "แก้ไขแล้ว"),
    ("9", "BGM startBGM()",
     "Method startBGM(String songName) ไม่ได้ใช้ Parameter songName "
     "แต่เล่น Media ที่ loadMediaPlayer() โหลดไว้ล่าสุดเสมอ "
     "แก้ไขโดยให้ startBGM() เรียก loadMediaPlayer(songName) ก่อนเล่น",
     "แก้ไขแล้ว"),
]

for bug in bugs:
    status = bug[3]
    row = tbl_bug.add_row()
    for i, text in enumerate(bug):
        cell = row.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        set_font(run, size=13)
    # Color the status cell
    status_cell = row.cells[3]
    tc = status_cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'C6EFCE' if status == 'แก้ไขแล้ว' else 'FFEB9C')
    tcPr.append(shd)
add_paragraph()

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
#  8. สรุป
# ═══════════════════════════════════════════════════════════════════════════
add_heading1("8. สรุปและข้อเสนอแนะ")

add_heading2("8.1 สรุปผลการพัฒนา")
add_body(
    "โปรแกรม Final Anime Survivor ได้พัฒนาเกมแนว Survivor ที่มีความซับซ้อนสูง "
    "ประกอบด้วยระบบอาวุธ 7 ชนิดพร้อม Evolved Form, Accessory 6 ชนิด, ระบบ Level Up, "
    "Chest, Backpack, Tile Map, Collision Detection และ VFX "
    "โครงสร้างโปรแกรมออกแบบตามหลัก OOP โดยใช้ Inheritance, Interface, Polymorphism "
    "และ Encapsulation อย่างครบถ้วนและสมเหตุสมผล"
)

add_heading2("8.2 สิ่งที่ควรพัฒนาเพิ่มเติม")
add_bullet("เพิ่ม Map หลายแบบและ Enemy หลากหลายกว่านี้ เพื่อเพิ่มความหลากหลายในการเล่น")
add_bullet("เพิ่มระบบ Save/Load Progress ให้ผู้เล่นสามารถเล่นต่อจากเดิมได้")
add_bullet("พัฒนา Achievement และ Collection System ให้สมบูรณ์ยิ่งขึ้น")
add_bullet("แก้ไข Shrine Sprite ให้แสดงภาพที่ถูกต้องและลบไฟล์ Duplicate ออก")
add_bullet("ปรับเงื่อนไขการใช้ Soda ให้ใช้ได้แม้ HP เต็ม เพื่อความสะดวกของผู้เล่น")

add_paragraph()
add_paragraph(
    "โดยรวมแล้วโปรแกรมสำเร็จตามวัตถุประสงค์หลักในการสร้างเกม Survivor ที่เล่นได้จริง "
    "มีกราฟิก เสียง และ Gameplay Loop ที่ครบถ้วน พร้อมการประยุกต์ใช้แนวคิด OOP ในระดับสูง",
    italic=True, size=15
)

# ─── Save ───────────────────────────────────────────────────────────────────
output_path = "/home/user/FinalAnimeSurvivor/report.docx"
doc.save(output_path)
print(f"Saved: {output_path}")
